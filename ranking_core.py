from __future__ import annotations

import io
import os
from typing import Any

import pandas as pd

DEFAULT_PILOTO_IMG = os.path.join("images", "Pilotodesc.png")

PILOTO_IMAGE_ALIASES: dict[str, str] = {
    "Helio Castroneves": "Helio Castroneves",
    "Hélio Castroneves": "Helio Castroneves",
}

# Cores do app (CSS) → (fill_hex, font_hex) para Excel.
PILOTO_CORES_EXCEL: dict[str, tuple[str, str]] = {
    "Gabriel Casagrande": ("800080", "FFFFFF"),  # purple
    "Lucas Foresti": ("808080", "FFFFFF"),  # gray
    "Cesar Ramos": ("FFFF00", "000000"),  # yellow
    "Thiago Camilo": ("FF0000", "FFFFFF"),  # red
    "Helio Castroneves": ("008000", "FFFFFF"),  # green
    "Renan Guerra": ("89CFF0", "000000"),  # baby blue
}

HEADER_FILL_HEX = "1A202C"
HEADER_FONT_HEX = "F7FAFC"
DEFAULT_ROW_FILL_HEX = "2D3748"
DEFAULT_ROW_FONT_HEX = "E2E8F0"


def exportar_ranking_excel(
    df: pd.DataFrame,
    *,
    sheet_name: str = "Ranking",
    titulo: str | None = None,
) -> bytes:
    """Gera .xlsx com as cores padrão dos pilotos (mesmo visual da tabela no app)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    thin = Border(
        left=Side(style="thin", color="4A5568"),
        right=Side(style="thin", color="4A5568"),
        top=Side(style="thin", color="4A5568"),
        bottom=Side(style="thin", color="4A5568"),
    )
    center = Alignment(horizontal="center", vertical="center")

    start_row = 1
    if titulo:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
        cell = ws.cell(1, 1, titulo)
        cell.font = Font(bold=True, size=14, color=HEADER_FONT_HEX)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
        cell.alignment = center
        start_row = 3

    header_row = start_row
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(header_row, col_idx, col_name)
        cell.font = Font(bold=True, color=HEADER_FONT_HEX)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
        cell.alignment = center
        cell.border = thin

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = header_row + 1 + row_offset
        piloto = str(row.get("Piloto", ""))
        fill_hex, font_hex = PILOTO_CORES_EXCEL.get(
            piloto, (DEFAULT_ROW_FILL_HEX, DEFAULT_ROW_FONT_HEX)
        )
        fill = PatternFill("solid", fgColor=fill_hex)
        font = Font(bold=piloto in PILOTO_CORES_EXCEL, color=font_hex)

        for col_idx, col_name in enumerate(df.columns, start=1):
            valor: Any = row[col_name]
            if pd.isna(valor):
                valor = ""
            elif hasattr(valor, "item"):
                try:
                    valor = valor.item()
                except (ValueError, AttributeError):
                    pass
            cell = ws.cell(excel_row, col_idx, valor)
            cell.fill = fill
            cell.font = font
            cell.alignment = center
            cell.border = thin

    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(
            len(str(col_name)),
            *(len(str(v)) for v in df[col_name].tolist()),
            8,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 36)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def strip_cell_header(header) -> str:
    if header is None:
        return ""
    return str(header).strip()


def normalizar_pdf_stockcar_2026(df_raw: pd.DataFrame) -> pd.DataFrame:
    """
    Layout típico do PDF Stock Car (pilotos): após Modelo repete-se pole («.») +
    corrida + corrida, e à direita (opcional) Descarte e Soma.
    """
    df = df_raw.copy()
    meta = ["Posição", "Numeral", "Piloto", "Equipe", "Modelo"]
    raw_cols = list(df.columns)
    n = len(raw_cols)
    if n <= 5:
        return df

    final = [None] * n
    for i in range(5):
        final[i] = meta[i]

    j = n - 1
    trailers = []
    while j >= 5:
        h_raw = strip_cell_header(raw_cols[j])
        h = h_raw.lower()
        if not h_raw:
            if j == n - 1:
                trailers.append((j, "Soma"))
                j -= 1
                continue
            break
        if "descarte" in h:
            trailers.append((j, "Descarte"))
            j -= 1
            continue
        if "soma" in h or h == "total":
            trailers.append((j, "Soma"))
            j -= 1
            continue
        break

    core_idxs = list(range(5, j + 1))
    pole_i = 1
    race_i = 1
    for offset, idx_col in enumerate(core_idxs):
        if offset % 3 == 0:
            final[idx_col] = f"pole_{pole_i}"
            pole_i += 1
        else:
            final[idx_col] = str(race_i)
            race_i += 1

    for idx_col, label in sorted(trailers, key=lambda t: t[0]):
        final[idx_col] = label

    for i in range(n):
        if final[i] is None:
            final[i] = f"extra_{i}"

    df.columns = final
    return df


def _coluna_tem_resultado(serie: pd.Series) -> bool:
    for valor in serie:
        if pd.isna(valor):
            continue
        texto = str(valor).strip().upper()
        if texto in {"", ".", "NP"}:
            continue
        return True
    return False


def corridas_logicas_da_etapa(etapa: int) -> tuple[int, int]:
    """Etapa de fim de semana (1, 2, …) → par sprint/principal (1-2, 3-4, …)."""
    return (2 * etapa - 1, 2 * etapa)


def colunas_protegidas_descarte(
    *,
    endurance_etapas: list[int],
    proteger_ultimas_concluidas: bool,
    ultima_corrida: int,
    total_corridas_ano: int | None,
    coluna_fn,
) -> set[str]:
    """
    Colunas que o regulamento não permite descartar.
    2026: Endurance + duas últimas provas do ano (sprint e principal).
    2025: duas últimas corridas já realizadas.
    """
    protegidas: set[str] = set()
    for etapa in endurance_etapas:
        for n in corridas_logicas_da_etapa(etapa):
            if n <= ultima_corrida:
                protegidas.add(str(coluna_fn(n)))

    if proteger_ultimas_concluidas and ultima_corrida >= 2:
        protegidas.add(str(coluna_fn(ultima_corrida - 1)))
        protegidas.add(str(coluna_fn(ultima_corrida)))

    if total_corridas_ano and total_corridas_ano >= 2:
        for n in (total_corridas_ano - 1, total_corridas_ano):
            if n <= ultima_corrida:
                protegidas.add(str(coluna_fn(n)))

    return protegidas


def detectar_ultima_corrida(df: pd.DataFrame, fallback: int = 1) -> int:
    """Infere a última corrida com dado no PDF (colunas numéricas 1, 2, 3…)."""
    corridas = sorted((int(c) for c in df.columns if str(c).isdigit()), reverse=True)
    for n in corridas:
        col = str(n)
        if col in df.columns and _coluna_tem_resultado(df[col]):
            return n
    return fallback


def caminho_imagem_piloto(nome_piloto: str) -> str:
    nome = PILOTO_IMAGE_ALIASES.get(nome_piloto, nome_piloto)
    path = os.path.join("images", f"{nome}.png")
    return path if os.path.isfile(path) else DEFAULT_PILOTO_IMG


def _parece_token_equipe(token: str) -> bool:
    t = token.strip()
    if not t or t.lower() in {"jr", "filho", "de", "da", "do", "di"}:
        return False
    if t.isupper() and len(t) <= 10:
        return True
    return len(t) <= 3 and t.isalpha() and t[0].isupper()


def _piloto_qualifying_de_tokens(tokens: list[str]) -> str:
    """Extrai só o nome do piloto (sem equipe) a partir dos tokens após posição/numeral."""
    if not tokens:
        return ""
    if len(tokens) == 1:
        partes = tokens[0].split()
        nome: list[str] = []
        for p in partes:
            if len(nome) >= 2 and _parece_token_equipe(p):
                break
            if len(nome) >= 3:
                break
            nome.append(p)
        return " ".join(nome).title()

    partes: list[str] = []
    for t in tokens:
        if len(partes) >= 2 and _parece_token_equipe(t):
            break
        if len(partes) >= 3:
            break
        partes.append(t)
    return " ".join(partes).title()


def normalizar_pilotos_qualifying(
    df_qualifying: pd.DataFrame, df_campeonato: pd.DataFrame
) -> pd.DataFrame:
    """Alinha nomes do qualifying aos da tabela oficial do campeonato (por numeral)."""
    out = df_qualifying.copy()
    if "Numeral" not in out.columns or "Numeral" not in df_campeonato.columns:
        return out

    mapa = (
        df_campeonato.assign(Numeral=pd.to_numeric(df_campeonato["Numeral"], errors="coerce"))
        .dropna(subset=["Numeral"])
        .drop_duplicates("Numeral", keep="first")
        .set_index("Numeral")["Piloto"]
    )
    nums = pd.to_numeric(out["Numeral"], errors="coerce")
    oficial = nums.map(mapa)
    out["Piloto"] = oficial.fillna(out["Piloto"])
    return out


def extrair_qualifying_pdf(arquivo_pdf: str) -> pd.DataFrame | None:
    """Extrai qualifying: tenta tabela estruturada; fallback para texto."""
    import pdfplumber

    dados: list[tuple[str, str, str]] = []
    try:
        with pdfplumber.open(arquivo_pdf) as pdf:
            pagina = pdf.pages[0]
            tabelas = pagina.extract_tables()
            if tabelas and tabelas[0]:
                for linha in tabelas[0]:
                    if not linha or not linha[0]:
                        continue
                    pos = str(linha[0]).strip()
                    if not pos.isdigit():
                        continue
                    no = str(linha[1]).strip() if len(linha) > 1 else ""
                    tokens = [
                        str(c).strip() for c in linha[2:] if c and str(c).strip()
                    ]
                    piloto = _piloto_qualifying_de_tokens(tokens)
                    if piloto:
                        dados.append((pos, no, piloto))
                if dados:
                    return pd.DataFrame(dados, columns=["Posição", "Numeral", "Piloto"])

            texto = pagina.extract_text() or ""
            for linha in texto.split("\n"):
                colunas = linha.split()
                if colunas and colunas[0].isdigit() and len(colunas) >= 3:
                    pos, no = colunas[0], colunas[1]
                    name = _piloto_qualifying_de_tokens(colunas[2:])
                    if name:
                        dados.append((pos, no, name))
    except OSError:
        return None

    if not dados:
        return None
    return pd.DataFrame(dados, columns=["Posição", "Numeral", "Piloto"])
